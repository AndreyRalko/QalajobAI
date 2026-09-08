"""Generate QalaJob AI project passport (.docx) from template structure."""

from __future__ import annotations

from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Install python-docx: pip install python-docx") from exc


def set_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)


def add_label_block(doc: Document, label: str, lines: list[str | tuple[str, bool]]) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: ")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    first = True
    for item in lines:
        if isinstance(item, tuple):
            text, bullet = item
            if bullet:
                bp = doc.add_paragraph(text, style="List Bullet")
                for r in bp.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
            else:
                np = doc.add_paragraph(text)
                for r in np.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
        elif item == "":
            doc.add_paragraph("")
        elif first:
            cont = p.add_run(item)
            cont.font.name = "Times New Roman"
            cont.font.size = Pt(12)
            first = False
        else:
            np = doc.add_paragraph(item)
            for r in np.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)


def build_document() -> Document:
    doc = Document()
    set_normal_style(doc)

    title = doc.add_paragraph()
    title_run = title.add_run("ПАСПОРТ КЕЙСА ИИ-проекта")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    add_label_block(
        doc,
        "ОРГАНИЗАЦИЯ",
        ['НАО «Кокшетауский университет имени Шокана Уалиханова»'],
    )
    doc.add_paragraph("")

    add_label_block(
        doc,
        "НАЗВАНИЕ КЕЙСА",
        ["QalaJob AI — платформа карьерного сопровождения студентов с ИИ"],
    )
    doc.add_paragraph("")

    add_label_block(
        doc,
        "ОПИСАНИЕ ПРОБЛЕМЫ / ЗАДАЧИ",
        [
            "Студенты вуза при выходе на рынок труда сталкиваются с рядом системных проблем:",
            ("отсутствие навыков составления профессионального резюме и сопроводительных писем;", True),
            ("сложность самостоятельного поиска релевантных вакансий и адаптации профиля под требования работодателей;", True),
            ("недостаточная подготовка к собеседованиям и отсутствие обратной связи по ответам;", True),
            ("разрыв между академическим профилем студента (данные LMS Platonus) и карьерными инструментами;", True),
            ("высокая нагрузка на сотрудников карьерного центра при индивидуальном сопровождении.", True),
            "",
            "Задачи проекта:",
            ("создать единый цифровой кабинет карьерного сопровождения студента;", True),
            ("автоматизировать сбор и использование учебных данных из LMS Platonus;", True),
            ("внедрить AI-помощника для резюме, писем и репетиции собеседований;", True),
            ("обеспечить подбор вакансий с hh.kz на основе профиля и интересов студента;", True),
            ("снизить время подготовки студента к трудоустройству;", True),
            ("повысить качество и прозрачность карьерного сопровождения в университете.", True),
        ],
    )
    doc.add_paragraph("")

    add_label_block(
        doc,
        "ПРОЦЕССЫ / ФУНКЦИИ ДЛЯ ВНЕДРЕНИЯ ИИ",
        [
            "Проект представляет собой веб-платформу (Next.js + Django REST) с гибридной AI-архитектурой: локальная модель (Ollama/Qwen) и облачный OpenAI с автоматическим fallback.",
            "",
            "1. AI-ассистент резюме",
            ("диалоговый сбор информации о студенте и формирование черновика резюме;", True),
            ("улучшение и структурирование текста резюме;", True),
            ("экспорт резюме в PDF.", True),
            "",
            "2. AI-генерация сопроводительных писем",
            ("формирование письма под конкретную вакансию и компанию;", True),
            ("использование данных резюме и контекста вакансии.", True),
            "",
            "3. AI-подготовка к собеседованию",
            ("репетиция по выбранной вакансии;", True),
            ("mock-interview: один вопрос за раз с обратной связью.", True),
            "",
            "4. AI-подбор вакансий (hh.kz)",
            ("формирование поискового запроса на основе транскрипта, интересов и профиля;", True),
            ("ранжирование вакансий и рекомендации;", True),
            ("адаптация резюме под выбранную вакансию.", True),
            "",
            "5. Интеграция с LMS Platonus",
            ("ежедневная синхронизация студентов (ФИО, логин, пароль, StudentID) через SSH-туннель;", True),
            ("синхронизация академического транскрипта для персонализации AI-рекомендаций.", True),
            "",
            "6. Администрирование и мониторинг AI",
            ("журнал AI-запросов (provider, статус, feature);", True),
            ("управление расписанием LMS sync;", True),
            ("роли: студент, работодатель, администратор.", True),
        ],
    )
    doc.add_paragraph("")

    add_label_block(
        doc,
        "ОЖИДАЕМЫЙ ЭКОНОМИЧЕСКИЙ ЭФФЕКТ (в том числе % от закладываемого бюджета), ПОКАЗАТЕЛИ ДЛЯ ОЦЕНКИ ЭФФЕКТИВНОСТИ (KPI)",
        [
            "Экономический эффект:",
            "",
            "Снижение трудозатрат карьерного центра",
            ("автоматизация типовых консультаций (резюме, письма, собеседования);", True),
            ("снижение ручной работы консультантов до 50–60%.", True),
            "",
            "Повышение качества подготовки студентов",
            ("стандартизация карьерных документов;", True),
            ("персонализированные рекомендации на основе учебного профиля.", True),
            "",
            "Ускорение трудоустройства выпускников",
            ("сокращение времени подготовки к отклику на вакансию;", True),
            ("рост доли релевантных откликов.", True),
            "",
            "Оптимизация затрат на AI",
            ("гибридная модель: локальный LLM для типовых задач, OpenAI — для сложных;", True),
            ("снижение расходов на облачные API до 40–50% по сравнению с полностью облачной схемой.", True),
            "",
            "KPI:",
            ("доля студентов, создавших резюме через платформу: ≥70%;", True),
            ("сокращение времени подготовки резюме: –60%;", True),
            ("доля автоматически синхронизированных профилей из Platonus: ≥95%;", True),
            ("точность релевантности рекомендованных вакансий (оценка пользователя): ≥75%;", True),
            ("уровень использования AI-кабинета активными студентами: ≥60%;", True),
            ("снижение обращений в карьерный центр по типовым вопросам: –50%.", True),
        ],
    )
    doc.add_paragraph("")

    add_label_block(
        doc,
        "КАК ЭТО БУДЕТ РАБОТАТЬ (ОПИСАНИЕ РЕШЕНИЯ)",
        [
            "6.1. Аутентификация и профиль",
            ("студент входит логином и паролем из Platonus;", True),
            ("система подтягивает ФИО, StudentID и академический транскрипт;", True),
            ("формируется единый профиль кандидата для AI-модулей.", True),
            "",
            "6.2. Рабочее пространство студента",
            ("AI-чат для резюме, сопроводительного письма и собеседования;", True),
            ("просмотр и редактирование черновиков документов;", True),
            ("подбор вакансий с hh.kz и сохранение избранных.", True),
            "",
            "6.3. Гибридная AI-обработка",
            ("запрос направляется в локальную модель (Ollama/Qwen) или OpenAI по правилам маршрутизации;", True),
            ("при сбое локальной модели выполняется fallback на OpenAI;", True),
            ("все обращения логируются для аудита и аналитики.", True),
            "",
            "6.4. Синхронизация с Platonus",
            ("по расписанию открывается SSH-туннель к серверу LMS;", True),
            ("из MySQL загружаются студенты и транскрипты;", True),
            ("данные сохраняются в локальной БД платформы;", True),
            ("туннель закрывается после завершения синхронизации.", True),
            "",
            "6.5. Администрирование",
            ("дашборды для управления пользователями, вакансиями и подписками;", True),
            ("журнал LMS sync и AI logs;", True),
            ("настройка расписания автосинхронизации.", True),
        ],
    )
    doc.add_paragraph("")

    add_label_block(
        doc,
        "НЕОБХОДИМЫЕ ДАННЫЕ",
        [
            ("данные студентов из LMS Platonus (ФИО, Login, Password, StudentID);", True),
            ("академический транскрипт (оценки, дисциплины, курс, семестр);", True),
            ("профиль и интересы студента (резюме, навыки, предпочтения по вакансиям);", True),
            ("данные вакансий через API hh.kz;", True),
            ("журналы AI-запросов и LMS sync для мониторинга;", True),
            ("техническая инфраструктура: сервер LMS (SSH), Ollama, Redis (опционально для Celery).", True),
        ],
    )
    doc.add_paragraph("")

    add_label_block(
        doc,
        "ДОПОЛНИТЕЛЬНЫЕ КОММЕНТАРИИ",
        [
            "Технологический стек: Next.js 16, React 19, Django 5, DRF, JWT, SQLite/PostgreSQL, Ollama, OpenAI API, HeadHunter API.",
            "Интерфейс платформы поддерживает казахский, русский и английский языки.",
            "Проект развёрнут как монорепозиторий с отдельными сервисами frontend (порт 3000) и backend API (порт 8000).",
            "Для production рекомендуется PostgreSQL, HTTPS и ограничение доступа к SSH/MySQL LMS.",
        ],
    )
    doc.add_paragraph("")

    add_label_block(
        doc,
        "ОТВЕТСТВЕННЫЕ ЛИЦА",
        ["Ралко Андрей Аркадьевич — старший программист."],
    )
    doc.add_paragraph("")

    p = doc.add_paragraph()
    run = p.add_run("СОГЛАСОВАНИЕ:")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    headers = ["№", "Ф.И.О.", "Должность", "Подпись"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        for paragraph in cell.paragraphs:
            for r in paragraph.runs:
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)

    row = table.rows[1].cells
    row[0].text = "1"
    row[1].text = "Сабитова Диана Сайрановна"
    row[2].text = "И.О. Проректора по интернационализации и развитию инфраструктуры"
    row[3].text = ""
    for cell in row:
        for paragraph in cell.paragraphs:
            for r in paragraph.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)

    return doc


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    filename = "Паспорт кейса QalaJob AI.docx"
    targets = [
        root / "docs" / filename,
        Path(r"d:\documents\Паспорта ИИ") / filename,
    ]

    doc = build_document()
    for target in targets:
        target.parent.mkdir(parents=True, exist_ok=True)
        doc.save(target)
        print(f"Saved: {target}")


if __name__ == "__main__":
    main()
